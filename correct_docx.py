import docx
import sys

def correct_docx(input_path, output_path):
    doc = docx.Document(input_path)
    
    for p in doc.paragraphs:
        if "MongoDB" in p.text or "MONGODB" in p.text or "mongodb" in p.text:
            modified_text = p.text.replace("MongoDB", "PostgreSQL").replace("MONGODB", "POSTGRESQL").replace("mongodb", "postgresql")
            for run in p.runs:
                run.text = ""
            p.add_run(modified_text)

    doc.save(output_path)

if __name__ == '__main__':
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    correct_docx(input_file, output_file)
