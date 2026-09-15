import docx
import sys
import re

def replace_text_in_runs(runs, old_text, new_text):
    text = "".join([run.text for run in runs])
    if old_text in text:
        text = text.replace(old_text, new_text)
        for i, run in enumerate(runs):
            run.text = text if i == 0 else ""

def correct_docx(input_path, output_path):
    doc = docx.Document(input_path)
    
    # Process paragraphs
    for p in doc.paragraphs:
        replace_text_in_runs(p.runs, "MongoDB est retenu pour la gestion des données", "PostgreSQL et SQLite sont retenus pour la gestion des données")
        replace_text_in_runs(p.runs, "MongoDB : la base de données permet", "PostgreSQL et SQLite : ces bases de données permettent")
        replace_text_in_runs(p.runs, "La couche de données correspond à MongoDB", "La couche de données correspond à PostgreSQL")
        replace_text_in_runs(p.runs, "une base de données MongoDB", "une base de données PostgreSQL")
        replace_text_in_runs(p.runs, "[Source : MongoDB Documentation, 2026.]", "")
        replace_text_in_runs(p.runs, "[5] MONGODB.", "[5] POSTGRESQL.")
        replace_text_in_runs(p.runs, "https://www.mongodb.com/docs/", "https://www.postgresql.org/docs/")
        replace_text_in_runs(p.runs, "base de données documentaire", "base de données relationnelle")

    # Process tables (just in case)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_text_in_runs(p.runs, "MongoDB", "PostgreSQL")

    doc.save(output_path)

if __name__ == '__main__':
    correct_docx(sys.argv[1], sys.argv[2])
